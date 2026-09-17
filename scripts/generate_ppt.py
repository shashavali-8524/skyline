from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pathlib import Path
slides=[
('SkyResolve AI','Auditable airline-disruption resolution\nAssignment 3 · Customer-Facing Resolution Agent'),
('Customer journey & goals','Understand intent and emotion\nAsk only necessary questions\nRecommend or simulate the correct next action\nEscalate when authority is missing\nPreserve the conversation and action record'),
('Grounded data, strict boundary','3 supplied customers · 4 supplied segments · 5 policy rules\nExercise date: Wednesday, 23 September 2026\nNo invented flights, seats, hotels, payment instruments, or compensation\nSample conversations guide tone only'),
('Production-shaped architecture','Streamlit UI  →  Agent orchestrator  →  Structured AI classification\n↓ Customer + booking context\nDeterministic policy engine  →  Guarded mocked tools\n↓\nSQLite decisions, actions, escalation, audit events'),
('Agent orchestration flow','1. Retrieve supplied customer and disruption\n2. Classify intent, entities, sentiment\n3. Route legal/formal complaints immediately\n4. Verify every candidate action against policy\n5. Invoke only authorized simulated tools\n6. Generate a grounded response and trace'),
('AI layer with safe fallback','Live Groq model when configured\nStructured, validated intents + entities + sentiment\nCustomer text treated as untrusted data\nLLM never decides eligibility or invokes tools\nNo key / malformed output / timeout → labeled offline fallback'),
('Policy & guardrails','Typed ALLOW · DENY · ESCALATE · ASK decisions\nOriginal-method refunds only · ₹1,500 waiver authority cap\nDelayed-hours hotel only · no extra loyalty compensation\nNo inventory invention · ownership validation · idempotency\nExact 3h / 5h equality escalates because source is silent'),
('Three scenario outcomes','PRIYA · Cancellation choice; cash/different-method refund and upgrade escalate\nARVIND · 4h: meal + lounge; hotel denied\nMEHER · 6h: meal + six-hour hotel; full night denied; ₹2,000 waiver escalated\nGold/Platinum add priority rebooking, not compensation'),
('Traceability & quality','SQLite captures messages, classifications, decisions, tools, actions, escalations\nCorrelation ID links input → AI → rules → tools → response\nCSV audit export in UI\nPytest covers scenarios, authority, ambiguity, ownership, and inventory\nAll external actions clearly marked SIMULATED'),
('Demo, limits & path to production','Demo: select each PNR, run scenario prompt, inspect decisions and trace\nCurrent limits: supplied static data, mocked airline actions, no real inventory\nNext: authenticated customer session, airline APIs, human queue, monitoring\nValue: empathetic AI experience with deterministic operational control')]
prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
for i,(title,body) in enumerate(slides,1):
 s=prs.slides.add_slide(prs.slide_layouts[6]); bg=s.background.fill; bg.solid(); bg.fore_color.rgb=RGBColor(247,249,252)
 top=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,0,0,prs.slide_width,Inches(.18)); top.fill.solid(); top.fill.fore_color.rgb=RGBColor(37,99,235); top.line.fill.background()
 tb=s.shapes.add_textbox(Inches(.75),Inches(.55),Inches(11.8),Inches(.8)).text_frame; p=tb.paragraphs[0]; p.text=title; p.font.size=Pt(28); p.font.bold=True; p.font.color.rgb=RGBColor(16,42,86)
 box=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(.8),Inches(1.65),Inches(11.75),Inches(4.65)); box.fill.solid(); box.fill.fore_color.rgb=RGBColor(255,255,255); box.line.color.rgb=RGBColor(220,229,242)
 tf=box.text_frame; tf.clear(); tf.margin_left=Inches(.45); tf.margin_right=Inches(.45); tf.margin_top=Inches(.35)
 for j,line in enumerate(body.split('\n')):
  p=tf.paragraphs[0] if j==0 else tf.add_paragraph(); p.text=line; p.font.size=Pt(20 if i!=1 else 25); p.font.color.rgb=RGBColor(35,48,72); p.space_after=Pt(14); p.level=0
 foot=s.shapes.add_textbox(Inches(.75),Inches(6.85),Inches(11.8),Inches(.3)).text_frame.paragraphs[0]; foot.text=f'Assignment 3 | Airline Disruption Resolution Agent                                      {i}/10'; foot.font.size=Pt(10); foot.font.color.rgb=RGBColor(102,112,133); foot.alignment=PP_ALIGN.CENTER
Path('artifacts').mkdir(exist_ok=True); out='artifacts/Assignment_3_Airline_Resolution_Agent.pptx'; prs.save(out)
assert len(Presentation(out).slides)==10
print(out)
